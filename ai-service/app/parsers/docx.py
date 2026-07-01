"""DOCX 파서 — python-docx로 본문/표 텍스트 추출 (명세서 §5.2).

DOCX는 렌더 시점에 페이지가 정해져 고정 페이지가 없다 → 전체를 단일 페이지로 낸다.
문단 + 표 셀 텍스트를 순서대로 이어붙인다.
"""

from __future__ import annotations

import io

from docx import Document as DocxDocument

from app.parsers.base import Page, ParsedDoc, ParserError


def parse_docx(data: bytes) -> ParsedDoc:
    try:
        doc = DocxDocument(io.BytesIO(data))
    except Exception as exc:  # 손상/비DOCX 등
        raise ParserError(f"DOCX 열기 실패: {exc!r}") from exc

    parts: list[str] = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    parts.append(cell_text)

    text = "\n".join(parts)
    if not text.strip():
        raise ParserError("DOCX에서 추출 가능한 텍스트가 없다.")
    return ParsedDoc(pages=[Page(page_no=1, text=text)], page_count=1)
