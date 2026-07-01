"""멀티포맷 파서(EPUB/TXT/DOCX) + 디스패처 테스트 — 픽스처는 코드로 생성."""

from __future__ import annotations

import io
import os
import tempfile

import pytest
from docx import Document as DocxDocument
from ebooklib import epub

from app.parsers import parse_document, supported_formats
from app.parsers.base import ParserError, UnsupportedFormatError
from app.parsers.docx import parse_docx
from app.parsers.epub import parse_epub
from app.parsers.txt import parse_txt


# ── 픽스처 생성 ──
def make_docx(paras: list[str], table: list[list[str]] | None = None) -> bytes:
    d = DocxDocument()
    for p in paras:
        d.add_paragraph(p)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for i, row in enumerate(table):
            for j, val in enumerate(row):
                t.rows[i].cells[j].text = val
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_epub(chapters: list[tuple[str, str]]) -> bytes:
    book = epub.EpubBook()
    book.set_identifier("id-test")
    book.set_title("Test Book")
    book.set_language("en")
    items = []
    for i, (title, html) in enumerate(chapters):
        c = epub.EpubHtml(title=title, file_name=f"c{i}.xhtml", lang="en")
        c.content = html
        book.add_item(c)
        items.append(c)
    book.toc = tuple(items)
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = ["nav", *items]
    fd, path = tempfile.mkstemp(suffix=".epub")
    os.close(fd)
    try:
        epub.write_epub(path, book)
        with open(path, "rb") as f:
            return f.read()
    finally:
        os.unlink(path)


# ── TXT ──
def test_txt_happy_single_page():
    doc = parse_txt("첫 줄\n둘째 줄".encode())
    assert doc.page_count == 1
    assert doc.pages[0].page_no == 1
    assert "둘째 줄" in doc.pages[0].text


def test_txt_cp949_decoding():
    doc = parse_txt("한글 인코딩".encode("cp949"))
    assert "한글" in doc.pages[0].text


def test_txt_empty_raises():
    with pytest.raises(ParserError):
        parse_txt(b"   \n  ")


# ── DOCX ──
def test_docx_paragraphs_and_table():
    data = make_docx(["문단 하나", "문단 둘"], table=[["셀A", "셀B"]])
    doc = parse_docx(data)
    assert doc.page_count == 1
    assert "문단 하나" in doc.pages[0].text
    assert "셀A" in doc.pages[0].text  # 표 셀도 포함


def test_docx_empty_raises():
    with pytest.raises(ParserError):
        parse_docx(make_docx([]))


def test_docx_corrupt_raises():
    with pytest.raises(ParserError):
        parse_docx(b"not a docx")


# ── EPUB ──
def test_epub_chapters_to_pages():
    data = make_epub([
        ("Ch1", "<html><body><h1>Chapter One</h1><p>alpha</p></body></html>"),
        ("Ch2", "<html><body><p>beta gamma</p></body></html>"),
    ])
    doc = parse_epub(data)
    assert doc.page_count == 2  # nav 제외, 챕터 2개
    assert "alpha" in doc.pages[0].text
    assert "beta gamma" in doc.pages[1].text


def test_epub_corrupt_raises():
    with pytest.raises(ParserError):
        parse_epub(b"not an epub")


# ── 디스패처 ──
def test_dispatcher_routes_by_format():
    doc = parse_document("TXT", "안녕".encode())  # 대문자·공백 정규화
    assert doc.pages[0].text.strip() == "안녕"


def test_dispatcher_unsupported_format():
    with pytest.raises(UnsupportedFormatError):
        parse_document("hwp", b"x")


def test_supported_formats_includes_new():
    fmts = supported_formats()
    assert {"pdf", "epub", "txt", "docx"} <= set(fmts)
